import io
from pathlib import Path

import pypdfium2 as pdfium
import pytesseract
from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.config import settings

if settings.tesseract_cmd:
    pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd

MIN_PAGE_CHARS = 20  # below this, the page is treated as scanned


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    rendered = None
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if len(text.strip()) < MIN_PAGE_CHARS:
            rendered = rendered or pdfium.PdfDocument(data)
            image = rendered[i].render(scale=300 / 72).to_pil()
            text = pytesseract.image_to_string(image, lang="eng")
        pages.append(text)
    return "\n".join(pages)


def _extract_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    return data.decode("utf-8", errors="ignore")